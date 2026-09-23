"""Etapa 5 - Gera o painel HTML e a planilha Excel.

Entradas: dados/agenda.json, dados/cruzamento.json, dados/potenciais.json, config
Saídas  : saida/painel.html   (arquivo único; abre em qualquer navegador)
          saida/rede_evento.xlsx (abas: Resumo, 1 - Rede na agenda, 1b - A verificar, 2 - Potenciais contatos)
Opcional: config/minha_agenda.ics ou .csv marca no painel as sessões que o pesquisador vai assistir.

Prioridade das pessoas da REDE:
  A = na mesma sessão do pesquisador, ou Confirmado em papel de destaque;
  B = demais Confirmados;  C = Prováveis (conferir).
Prioridade dos POTENCIAIS vem da etapa 4.
"""
import json
import os
import re
import sys
import urllib.parse

sys.path.insert(0, os.path.dirname(__file__))
from comum import RAIZ, carregar_config, chave_data, chave_nome, instituicao_curta, instituicoes, ler_json  # noqa: E402

DESTAQUE = {"Palestrante", "Painelista", "Proponente", "Moderador", "Coordenador", "Coordenador/Debatedor",
            "Debatedor", "Editor", "Mentor", "Ministrante", "Diretoria ANPAD"}
SAIDA = os.path.join(RAIZ, "saida")


def evento(r, pid):
    return dict(i=pid, d=r["data"], t=r["hora"], rm=r["sala"], s=r["sessao"] or "", dv=r["divisao"],
                ro=r["papel"], w=(r["titulo"] or "")[:160], c=r["codigo"] or "")


def ler_agenda_pessoal(cfg, agenda):
    """Sessões que o pesquisador marcou para assistir (opcional).

    Aceita config/minha_agenda.ics (exportado do Google Agenda, Outlook etc.) ou
    config/minha_agenda.csv com colunas data;hora;sala (ex.: 24/09/2026;10:45;Sala B44 - 2º Andar - Bloco B).
    Cada evento é associado à sessão da programação com mesma data, horário de início e sala;
    eventos sem correspondência (ex.: cerimônias fora da programação) só entram se o título casar com
    evento.filtro_agenda_pessoal (expressão regular, ex.: "ANPAD"); os demais compromissos são ignorados.
    """
    import csv
    from datetime import datetime
    from zoneinfo import ZoneInfo
    fuso = ZoneInfo(cfg["evento"].get("fuso_horario", "America/Sao_Paulo"))
    datas = {r["data"] for r in agenda}
    filtro = cfg["evento"].get("filtro_agenda_pessoal", "")
    brutos = []
    ics = os.path.join(RAIZ, "config", "minha_agenda.ics")
    csvp = os.path.join(RAIZ, "config", "minha_agenda.csv")
    if os.path.exists(ics):
        texto = re.sub(r"\r?\n[ \t]", "", open(ics, encoding="utf-8").read())
        for bloco in texto.split("BEGIN:VEVENT")[1:]:
            campos = {}
            for linha in bloco.splitlines():
                if ":" in linha:
                    k, v = linha.split(":", 1)
                    campos[k.split(";")[0]] = v.strip()

            def hora(v):
                if not v or "T" not in v:
                    return None
                dt = datetime.strptime(v.rstrip("Z")[:15], "%Y%m%dT%H%M%S")
                if v.endswith("Z"):
                    dt = dt.replace(tzinfo=ZoneInfo("UTC")).astimezone(fuso)
                return dt
            ini, fim = hora(campos.get("DTSTART")), hora(campos.get("DTEND"))
            if ini:
                brutos.append(dict(d=ini.strftime("%d/%m/%Y"), h=ini.strftime("%H:%M"),
                                   hf=fim.strftime("%H:%M") if fim else "",
                                   rm=campos.get("LOCATION", "").replace("\\,", ",").strip(),
                                   titulo=campos.get("SUMMARY", "").replace("\\,", ",")))
    elif os.path.exists(csvp):
        with open(csvp, encoding="utf-8") as f:
            for l in csv.DictReader(f, delimiter=";"):
                brutos.append(dict(d=l["data"].strip(), h=l["hora"].strip()[:5], hf="", rm=l["sala"].strip(), titulo=l.get("titulo", "")))
    saida = []
    for b in brutos:
        if b["d"] not in datas:
            continue  # evento fora dos dias do congresso
        m = [r for r in agenda if r["data"] == b["d"] and (r["hora"] or "").startswith(b["h"])
             and b["rm"] and r["sala"] == b["rm"]]
        if m:
            saida.append(dict(d=m[0]["data"], t=m[0]["hora"], rm=m[0]["sala"], s=m[0]["sessao"] or "", dv=m[0]["divisao"]))
        elif filtro and re.search(filtro, b["titulo"], re.I):
            saida.append(dict(d=b["d"], t=f'{b["h"]} - {b["hf"]}' if b["hf"] else b["h"], rm=b["rm"], s=b["titulo"], dv=""))
    unicos = {(x["d"], x["t"], x["rm"]): x for x in saida}
    if unicos:
        print(f"[ok] agenda pessoal: {len(unicos)} sessões marcadas")
    return list(unicos.values())


def montar_dados(cfg):
    agenda = ler_json("agenda.json")
    cruz = ler_json("cruzamento.json")
    pot = ler_json("potenciais.json")
    eu = chave_nome(cfg.get("pesquisador", {}).get("nome_na_agenda", ""))
    aliases = cfg["evento"].get("aliases_instituicoes", {})
    minhas = [r for r in agenda if eu and chave_nome(r["nome"]) == eu]
    chaves_minhas = {(r["data"], r["hora"], r["sala"]) for r in minhas}
    meus_codigos = {r["codigo"] for r in minhas if r["codigo"]}
    por_chave = {}
    for r in agenda:
        por_chave.setdefault(chave_nome(r["nome"]), []).append(r)

    P, E, pid = [], [], 0
    rede = {}
    for c in cruz:
        if c["status"] in ("Confirmado", "Provável"):
            rede.setdefault(c["chave_linkedin"], {"c": c, "agks": set()})["agks"].add(c["chave_agenda"])
    for k, v in rede.items():
        c = v["c"]
        itens = [r for a in v["agks"] for r in por_chave.get(a, [])]
        if not itens:
            continue
        mesma = any((r["data"], r["hora"], r["sala"]) in chaves_minhas for r in itens)
        papeis = {r["papel"] for r in itens}
        pr = "A" if mesma or (papeis & DESTAQUE and c["status"] == "Confirmado") else ("B" if c["status"] == "Confirmado" else "C")
        nota = ("Coautor(a) do seu trabalho" if any(r["codigo"] in meus_codigos for r in itens if r["codigo"])
                else "Na sua sessão" if mesma else "Papel de destaque" if papeis & DESTAQUE else "")
        P.append(dict(i=pid, n=c["nome_linkedin"], g="r", p=pr, st=c["status"], li=" · ".join(c["vinculos"]),
                      h=c["titulo"][:140], u=c["url"], af=instituicao_curta(itens[0]["afiliacao"]), note=nota,
                      ins=instituicoes([r["afiliacao"] for r in itens], aliases),
                      sg=(instituicoes([itens[0]["afiliacao"].split(") - (")[0]], aliases) or [""])[0]))
        E += [evento(r, pid) for r in itens]
        pid += 1
    for x in pot:
        P.append(dict(i=pid, n=x["nome"], g="p", p=x["prioridade"], st="", li="", h="", u=x.get("url", ""),
                      af=instituicao_curta(x["afiliacao"]), note="Na sua sessão" if x["mesma_sessao"] else "",
                      temas=x["temas"], ins=instituicoes([r["afiliacao"] for r in x["itens"]], aliases),
                      sg=(instituicoes([x["afiliacao"].split(" | ")[0].split(") - (")[0]], aliases) or [""])[0]))
        E += [evento(r, pid) for r in x["itens"]]
        pid += 1
    E = [e for e in E if e["d"]]
    minhas_sessoes = []
    vistos = set()
    for r in sorted(minhas, key=lambda r: chave_data(r["data"], r["hora"])):
        chave = (r["data"], r["hora"], r["sala"])
        if chave in vistos:
            continue
        vistos.add(chave)
        papeis = sorted({m["papel"] for m in minhas if (m["data"], m["hora"], m["sala"]) == chave})
        minhas_sessoes.append(dict(d=r["data"], t=r["hora"], rm=r["sala"], s=r["sessao"], dv=r["divisao"], papeis=papeis))
    ev = cfg["evento"]
    meta = dict(evento=ev["nome"], local=ev["local"], periodo=ev["periodo"], dias=ev["rotulos_dias"],
                divisoes=ev["divisoes"], minhas=minhas_sessoes,
                agendadas=ler_agenda_pessoal(cfg, agenda))
    return dict(meta=meta, P=P, E=E), agenda, cruz, pot


def gerar_painel(dados):
    tpl = open(os.path.join(RAIZ, "template", "painel_template.html"), encoding="utf-8").read()
    js = json.dumps(dados, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = tpl.replace("__DATA__", js).replace("__TITULO__", "Radar " + dados["meta"]["evento"])
    os.makedirs(SAIDA, exist_ok=True)
    caminho = os.path.join(SAIDA, "painel.html")
    open(caminho, "w", encoding="utf-8").write(html)
    print("[ok] saida/painel.html")


def gerar_planilha(dados, cruz, pot):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    HFILL = PatternFill("solid", fgColor="1F3864")
    BF = Font(name="Arial", size=10)
    LF = Font(name="Arial", size=10, color="0563C1", underline="single")
    wb = Workbook()

    def aba(titulo, cab, linhas, larg, nome_tab, col_link=None):
        ws = wb.create_sheet(titulo)
        for j, h in enumerate(cab, 1):
            c = ws.cell(1, j, h)
            c.font, c.fill = HF, HFILL
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, lin in enumerate(linhas, 2):
            for j, v in enumerate(lin, 1):
                c = ws.cell(i, j, v)
                c.font, c.alignment = BF, Alignment(wrap_text=True, vertical="top")
                if col_link == j and isinstance(v, str) and v.startswith("http"):
                    c.hyperlink, c.font = v, LF
        for j, w in enumerate(larg, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "C2"
        if linhas:
            t = Table(displayName=nome_tab, ref=f"A1:{get_column_letter(len(cab))}{len(linhas) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(t)

    byid = {p["i"]: p for p in dados["P"]}
    minhas = {(m["d"], m["t"], m["rm"]) for m in dados["meta"]["minhas"]}
    l1 = []
    for e in sorted(dados["E"], key=lambda e: chave_data(e["d"], e["t"])):
        p = byid[e["i"]]
        if p["g"] != "r":
            continue
        l1.append([p["p"], p["st"], p["n"], p["li"], p["h"], p["u"], e["ro"], p["af"], e["dv"], e["d"], e["t"], e["rm"], e["s"],
                   (e["c"] + " - " + e["w"]) if e["w"] else "", "Sim" if (e["d"], e["t"], e["rm"]) in minhas else "", p["note"], ""])
    cab1 = ["Prioridade", "Status", "Nome no LinkedIn", "Vínculo", "Título (LinkedIn)", "Perfil", "Papel", "Instituição", "Divisão",
            "Data", "Horário", "Sala", "Sessão", "Trabalho", "Na sua sessão?", "Observação", "Feito? / notas"]
    aba("1 - Rede na agenda", cab1, l1, [10, 12, 26, 20, 40, 30, 16, 30, 10, 11, 13, 24, 40, 50, 10, 24, 16], "Rede", 6)
    l1b = [[c["status"], c["nome_linkedin"], ", ".join(c["vinculos"]), c["titulo"][:200], c["url"], c["chave_agenda"]]
           for c in cruz if c["status"] == "A verificar"]
    aba("1b - A verificar", ["Status", "Nome no LinkedIn", "Vínculo", "Título", "Perfil", "Nome na agenda (normalizado)"],
        l1b, [12, 28, 20, 50, 34, 34], "Verificar", 5)
    l2 = []
    for x in pot:
        itens = sorted(x["itens"], key=lambda r: chave_data(r["data"], r["hora"]))
        busca = "https://www.google.com/search?q=" + urllib.parse.quote(f'site:linkedin.com/in "{x["nome"]}" {instituicao_curta(x["afiliacao"])}')
        l2.append([x["prioridade"], x["pontuacao"], x["nome"], instituicao_curta(x["afiliacao"]), ", ".join(x["papeis"]), "; ".join(x["temas"][:4]),
                   "; ".join(f'{r["data"][:5]} {r["hora"][:5]} ({r["papel"]}) {r["sessao"][:70] if r["sessao"] else ""}' for r in itens[:6]),
                   "Sim" if x["mesma_sessao"] else "", x.get("url") or busca, "Localizado" if x.get("url") else "Buscar", ""])
    aba("2 - Potenciais contatos", ["Prioridade", "Pontuação", "Nome", "Instituição", "Papéis", "Aderência temática", "Participações",
                                    "Na sua sessão?", "Perfil / busca", "Situação do perfil", "Feito? / notas"],
        l2, [10, 10, 28, 30, 24, 36, 70, 10, 36, 14, 16], "Potenciais", 9)
    ws = wb["Sheet"]
    ws.title = "Resumo"
    linhas = [("Indicador", "Valor"), ("Pessoas da rede na agenda (A+B+C)", sum(1 for p in dados["P"] if p["g"] == "r")),
              ("Potenciais contatos", '=ROWS(Potenciais[Nome])'), ("Potenciais - prioridade A", '=COUNTIF(Potenciais[Prioridade],"A")'),
              ("Potenciais - prioridade B", '=COUNTIF(Potenciais[Prioridade],"B")'), ("Potenciais - prioridade C", '=COUNTIF(Potenciais[Prioridade],"C")')]
    for i, (a, b) in enumerate(linhas, 1):
        ws.cell(i, 1, a).font = HF if i == 1 else BF
        ws.cell(i, 2, b).font = HF if i == 1 else BF
        if i == 1:
            ws.cell(i, 1).fill = ws.cell(i, 2).fill = HFILL
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 42, 14
    os.makedirs(SAIDA, exist_ok=True)
    wb.save(os.path.join(SAIDA, "rede_evento.xlsx"))
    print("[ok] saida/rede_evento.xlsx")


def main():
    cfg = carregar_config()
    dados, agenda, cruz, pot = montar_dados(cfg)
    gerar_painel(dados)
    gerar_planilha(dados, cruz, pot)


if __name__ == "__main__":
    main()
